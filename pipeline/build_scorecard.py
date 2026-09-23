"""Create the aggregate-only Rounds workbook from manifest.json."""
from pathlib import Path
import json,datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

BASE=Path(__file__).resolve().parents[1]/'public'/'downloads'
DATA=json.loads((BASE.parents[1]/'results'/'manifest.json').read_text(encoding='utf8'))
wb=Workbook(); summary=wb.active; summary.title='Scorecard'
source=wb.create_sheet('Measures'); notes=wb.create_sheet('Notes')
NAVY='18364B'; BLUE='EAF1F7'; TEXT='203443'; MUTED='536574'; LINE='D5E0E8'
font='Arial'
for ws in wb:
 ws.sheet_view.showGridLines=False
 ws.sheet_properties.pageSetUpPr.fitToPage=True
 ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A3
 ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
 ws.print_options.horizontalCentered=True
 ws.page_margins.left=.35; ws.page_margins.right=.35
 ws.page_margins.top=.45; ws.page_margins.bottom=.45
 ws.oddFooter.center.text='rounds | Synthetic data | Not for clinical use'
 ws.oddFooter.right.text='Page &P of &N'
 ws.sheet_view.zoomScale=90

def put(ws,row,col,value,bold=False,color=TEXT,size=10,wrap=False):
 c=ws.cell(row,col,value); c.font=Font(name=font,size=size,bold=bold,color=color)
 c.alignment=Alignment(vertical='center',horizontal='left',wrap_text=wrap)
 return c
def title(ws,text,last):
 put(ws,2,1,text,True,size=16); ws.row_dimensions[2].height=27
 for col in range(1,last+1): ws.cell(3,col).border=Border(bottom=Side(style='thin',color=NAVY))
def headers(ws,row,labels):
 for col,label in enumerate(labels,1):
  c=put(ws,row,col,label,True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor=NAVY)
  c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
  c.border=Border(right=Side(style='thin',color='FFFFFF'))
 ws.row_dimensions[row].height=30
def band(ws,row,cols):
 if row%2==0:
  for col in range(1,cols+1): ws.cell(row,col).fill=PatternFill('solid',fgColor='F2F6F9')
def numeric(c,fmt='#,##0'):
 c.number_format=fmt;c.alignment=Alignment(horizontal='right',vertical='center')

summary.sheet_properties.tabColor=NAVY
title(summary,'rounds quality scorecard',6)
put(summary,4,1,'Synthetic Massachusetts sample. As of '+DATA['provenance']['as_of']+'.')
put(summary,5,1,'Simplified documentation and data-coverage measures. No certified quality targets or benchmark thresholds.',color=MUTED)
headers(summary,7,['Measure','Department','Numerator','Denominator','Rate','Scope'])
source.freeze_panes='C8'; summary.freeze_panes='C8'
for i,m in enumerate(DATA['measures'],8):
 r=i
 vals=[m['name'],m['department'],f'=IF(ISNUMBER(Measures!C{r}),Measures!C{r},"n.a.")',f'=IF(ISNUMBER(Measures!D{r}),Measures!D{r},"n.a.")',f'=IF(OR(NOT(ISNUMBER(Measures!C{r})),NOT(ISNUMBER(Measures!D{r}))),"n.a.",IF(Measures!D{r}=0,"n.a.",Measures!C{r}/Measures!D{r}))','Lifetime index admissions' if m['id']=='return30' else ('Lifetime admission coverage' if m['id']=='followup-coverage' else 'Prior 365 days. Living population.')]
 for j,value in enumerate(vals,1): put(summary,r,j,value,wrap=j in [1,2,6])
 for j in [3,4]: numeric(summary.cell(r,j))
 numeric(summary.cell(r,5),'0.00%'); band(summary,r,6); summary.row_dimensions[r].height=35
put(summary,19,1,'Each denominator follows its own specification. Rates must not be summed or averaged across measures.',color=MUTED)
put(summary,20,1,'Definitions, exclusions and provenance are in Notes. Empty inputs and zero denominators display n.a.',color=MUTED)
put(summary,22,1,'Observed data checks',True,size=13)
put(summary,23,1,DATA['quality']['summary'],color=MUTED)
headers(summary,25,['Check','Category','Checked rows','Failed rows','Status','Scope'])
for r,check in enumerate(DATA['quality']['checks'],26):
 for j,key in enumerate(['name','category','checked','failed','status'],1): put(summary,r,j,check[key],wrap=j in [1,2])
 put(summary,r,6,'Observed-data check. No defect injection.',wrap=True)
 for j in [3,4]: numeric(summary.cell(r,j))
 summary.row_dimensions[r].height=30; band(summary,r,6)
summary.print_area='A1:F39'; summary.print_title_rows='1:7'
for col,width in zip('ABCDEF',[44,20,14,14,14,44]): summary.column_dimensions[col].width=width

title(source,'Measure source values',7)
put(source,4,1,'Source: generated manifest.json. Official Synthea build '+DATA['provenance']['version']+'. Seed '+str(DATA['provenance']['seed'])+'.')
put(source,5,1,'Each row represents one aggregate measure. Original numerators, denominators and reported rates are preserved.',color=MUTED)
headers(source,7,['Measure ID','Measure','Numerator','Denominator','Source rate','Unit','Department'])
for r,m in enumerate(DATA['measures'],8):
 values=[m['id'],m['name'],m['numerator'],m['denominator'],m['value']/100 if m['value'] is not None else None,m['unit'],m['department']]
 for j,value in enumerate(values,1): put(source,r,j,value,wrap=j in [1,2,7])
 for j in [3,4]: numeric(source.cell(r,j))
 numeric(source.cell(r,5),'0.00%'); band(source,r,7); source.row_dimensions[r].height=35
table=Table(displayName='MeasureSource',ref='A7:G17'); table.tableStyleInfo=TableStyleInfo(name='TableStyleLight9',showRowStripes=False); source.add_table(table)
source.auto_filter.ref='A7:G17'; source.print_area='A1:G19'; source.print_title_rows='1:7'
for col,width in zip('ABCDEFG',[23,44,14,14,16,10,22]): source.column_dimensions[col].width=width

title(notes,'Definitions and provenance',2)
put(notes,4,1,'Required statement',True)
put(notes,4,2,DATA['statement'],wrap=True); notes.row_dimensions[4].height=49
provenance=DATA['provenance']
facts=[('Population',provenance['population']+'; '+str(provenance['patients'])+' patients; '+str(provenance['encounters'])+' source encounters.'),('Source',provenance['source']),('Generation',f"Seed {provenance['seed']}. Build {provenance['version']}. Reference and simulation end {provenance['as_of']}."),('Calendar',provenance['date_basis']),('Missing values','Unavailable source values remain blank in Measures. Scorecard formulas return n.a. for missing required counts or a zero denominator. A real zero numerator produces a zero rate.'),('Interpretation','These are simplified, noncertified documentation or data-coverage measures. No benchmark, target, clinical recommendation or comparison against real hospitals is implied.'),('Recalculation','Scorecard numerators, denominators and rates link to Measures. The rate is numerator divided by denominator. Source rate preserves the rounded manifest value.'),('Data checks',DATA['quality']['summary']),('Terminology','Source-code presence is a completeness measure, not proof of standard OMOP concept mapping.'),('Privacy','This workbook contains aggregate values only. It contains no patient-level records, pseudonyms or identifiers.'),('Synthea source','https://github.com/synthetichealth/synthea/releases/tag/master-branch-latest')]
for r,(label,value) in enumerate(facts,6):
 put(notes,r,1,label,True); put(notes,r,2,value,wrap=True); notes.row_dimensions[r].height=34 if len(value)>130 else 25
notes.cell(16,2).hyperlink=notes.cell(16,2).value
put(notes,18,1,'Measure definitions',True,size=13)
for idx,m in enumerate(DATA['measures']):
 r=20+idx*3
 put(notes,r,1,m['name'],True,wrap=True); put(notes,r,2,m['definition'],wrap=True)
 notes.row_dimensions[r].height=50
 put(notes,r+1,1,'Exclusions',color=MUTED); put(notes,r+1,2,m['exclusions'],wrap=True); notes.row_dimensions[r+1].height=35
notes.column_dimensions['A'].width=39;notes.column_dimensions['B'].width=121
notes.freeze_panes='B6';notes.print_area='A1:B49';notes.print_title_rows='1:3'
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
wb.active=0
out=BASE/'rounds-quality-scorecard.xlsx';wb.save(out)
print(str(out))
